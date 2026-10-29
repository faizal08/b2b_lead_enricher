import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_custom_excel(records, company_name):
    """
    Formats extracted records into an Excel workbook matching the Master_Database template.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master_Database"
    
    headers = [
        'Record id', 'Company Name', 'Company type', 'Primary Industry', 'Country', 'City', 
        'Primary phone', 'Whatsapp number', 'Website URL', 'General Email', 'Employee count', 
        'DM Full name', 'DM Title', 'DM Authority Level', 'DM Direct Email', 'DM Direct phone', 
        'DM Linkedin URL', 'HQ Address', 'Company Linkedin'
    ]
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=9)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        
    for idx, row_data in enumerate(records, start=1):
        record_id = f"REC-2026-{idx:05d}"
        row = [
            record_id,
            row_data.get('company', ''),
            'Private Corporate Entity',
            'Technology / Real Estate',
            row_data.get('country', ''),
            row_data.get('city', ''),
            '', 
            '', 
            f"https://{row_data.get('domain', '')}",
            f"info@{row_data.get('domain', '')}",
            row_data.get('employee_count', ''),
            f"{row_data.get('first_name', '')} {row_data.get('last_name', '')}",
            row_data.get('title', ''),
            row_data.get('authority_level', ''),
            row_data.get('email', ''),
            row_data.get('phone', ''),
            row_data.get('linkedin_url', ''),
            f"{row_data.get('city', '')}, {row_data.get('country', '')}",
            f"https://www.linkedin.com/company/{row_data.get('company', '').lower().replace(' ', '')}"
        ]
        ws.append(row)
        
    for r in range(2, len(records) + 2):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = data_font
            cell.border = thin_border
            
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    safe_company_name = company_name.replace(' ', '_') if company_name else "Target"
    output_filename = f"Enriched_Leads_{safe_company_name}.xlsx"
    wb.save(output_filename)
    return output_filename